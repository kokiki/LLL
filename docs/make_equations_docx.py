"""
make_equations_docx.py -- CSI 예측 수식 정리 Word 문서 생성 -> docs/BSC-UNet_수식정리.docx
  각 수식: Word 수식(OMML, latex2mathml -> Office MML2OMML.XSL 변환) + 아래에 LaTeX 코드(PPT 수식 편집기 붙여넣기용)
  기준 구현: generate_ntn_csi_v9b_fdd_dl.m, data.py, csi_models.py, baselines.py, train_beam.py, evaluate.py
            (확장: generate_ntn_csi_v10_fdd_dl.m)
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from latex2mathml.converter import convert as tex2mml
from lxml import etree

XSL = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "BSC-UNet_수식정리.docx")
_xslt = etree.XSLT(etree.parse(XSL))
FAILED = []

# ------------------------------------------------------------------ 문서 내용
# ("h1"/"h2", 제목) | ("p", 문단) | ("eq", LaTeX) | ("impl", 구현 위치) | ("table", [헤더], [[행]...])
C = []
H1 = lambda t: C.append(("h1", t)); H2 = lambda t: C.append(("h2", t))
P = lambda t: C.append(("p", t)); EQ = lambda t: C.append(("eq", t)); IMPL = lambda t: C.append(("impl", t))
TB = lambda hdr, rows: C.append(("table", hdr, rows))

P("모든 수식은 실제 구현과 기호·정의가 일치하도록 작성했다. 각 수식 아래 회색 글씨는 PowerPoint 수식 편집기"
  "(삽입 → 수식 → LaTeX 형식)에 그대로 붙여 넣을 수 있는 LaTeX 코드이고, ‘구현’은 대응하는 파일·함수다. "
  "주 결과는 v9b 데이터(위성 기준 방위각 0° 고정), 확장 결과는 v10 데이터(구면 궤도 기하, 10절)이다.")

H1("1. 표기법")
TB(["기호", "의미"], [
    ["M = L = 16", "위성 UPA 행·열 소자 수 (Nt = 256), 단말 안테나 1개"],
    ["m, n", "배열 행·열 인덱스 (0, …, 15)"],
    ["q", "부반송파 인덱스. 저장 0…299, 사용 𝒬 = {0, 25, …, 275} (12개)"],
    ["i, tᵢ = iΔt", "관측(시간) 인덱스, Δt = 1 ms, 에피소드당 i = 0…149"],
    ["H_q(i) ∈ ℂ^{16×16}, h_q(i) = vec(H_q(i))", "부반송파 q, 관측 i의 CSI (행 우선 펼침: 원소 16m+n)"],
    ["T_p = 20, k", "history 길이, 예측 horizon (k ∈ {1, 2, 4, 8, 12, 16, 24, 32} ms)"],
    ["T_f", "윈도우 미래 길이 (k ≤ 16이면 16, k = 24, 32이면 k)"],
    ["h̃, h̄", "잡음 관측 / 인과적 정규화·Doppler 잔차 회전 제거 후 신호"],
    ["a(α, β)", "평면파 조향 벡터 [a]_mn = e^{j(αm+βn)}/16 (단위 노름)"],
    ["s(τ)", "빔 계수 s(τ) = aᴴ h̃(τ)"],
])

H1("2. 시스템 및 시나리오 파라미터")
TB(["항목", "값"], [
    ["반송파 / 이중화 / 링크", "f_c = 20 GHz (Ka-band), FDD, 하향링크 (위성 → 단말)"],
    ["위성 궤도", "원궤도, 고도 h = 550 km, 속도 v_sat = 7.6 km/s, R_e = 6371 km"],
    ["위성 배열", "16×16 UPA, 단일 편파 (nrCDLChannel 기본 소자 간격·소자 패턴)"],
    ["뉴머롤로지", "SCS 60 kHz, 25 RB (300 부반송파), N_FFT = 512, normal CP"],
    ["관측", "Δt = 1 ms (4 슬롯마다 1 슬롯, 슬롯 마지막 OFDM 심볼의 주파수 응답)"],
    ["채널 모델", "3GPP TR 38.811 NTN-CDL, Urban (LOS: NTN-CDL-C, NLOS: NTN-CDL-A)"],
    ["단말 속도", "v_UE = 0.25 km/h → f_D,UE = v_UE·f_c/c = 4.63 Hz"],
    ["에피소드", "400개 × 150 ms, 에피소드 내 대규모 파라미터 고정"],
])
IMPL("generate_ntn_csi_v9b_fdd_dl.m  §1 파라미터")

H1("3. 채널 데이터 생성: 3GPP TR 38.811 NTN-CDL")
H2("3.1 고각 프로파일과 고각 구간 분할")
P("에피소드 e = 1…400의 위성 고각은 하나의 통과를 본뜬 정현 스윕으로 정한다.")
EQ(r"\varepsilon_e=\varepsilon_{\min}+(\varepsilon_{\max}-\varepsilon_{\min})\sin\left(\pi\frac{e-1}{399}\right),\qquad \varepsilon_{\min}=10^\circ,\ \varepsilon_{\max}=90^\circ")
P("고각 구간으로 split을 정한다 (에피소드 단위 배정, 한 에피소드의 모든 윈도우는 같은 split). "
  "test: [12.5, 14.5), [20.5, 24.5), [40.5, 44.5), [60.5, 64.5), [81, 83)°, val: [30, 35), [50, 55), [70, 75)°, "
  "guard: test 구간 양쪽 1.5–2.5° (train–test 완충, 학습 미사용), train: 나머지. "
  "결과: train 186 / val 60 / test 68 / guard 86 에피소드 (NLOS 27 / 16 / 25 / 33).")
IMPL("generateElevationProfile, episodeSplit (testBandsDeg, valBandsDeg, testGuardsDeg)")

H2("3.2 LOS/NLOS 상태와 대규모 파라미터 (에피소드마다 1회 샘플링)")
P("LOS 확률 P_LOS(ε)는 TR 38.811 Table 6.6.1-1 (Urban)을 10° 격자에서 선형보간한다.")
EQ(r"\mathrm{NLOS}\iff U>P_{\mathrm{LOS}}(\varepsilon_e),\qquad U\sim\mathcal{U}[0,1]")
P("지연 확산 (Table 6.7.2-3b / 4b, Ka-band):")
EQ(r"\log_{10}\mathrm{DS}\sim\mathcal{N}\left(\mu_{\lg\mathrm{DS}}(\varepsilon_e),\,\sigma_{\lg\mathrm{DS}}^2(\varepsilon_e)\right),\qquad \mathrm{DS}\in[1\,\mathrm{ns},\,1\,\mu\mathrm{s}]")
P("K-factor (LOS만, Table 6.7.2-3b, Ka-band):")
EQ(r"K_{\mathrm{dB}}\sim\mathcal{N}\left(\mu_K(\varepsilon_e),\,\sigma_K^2(\varepsilon_e)\right)")
P("Shadow fading / NLOS clutter loss (Table 6.6.2-2, Ka-band): LOS는 μ_SF = 0, σ_SF = 4 dB, "
  "NLOS는 μ_SF = CL(ε) (44.3 dB@10° ~ 32.9 dB@90°), σ_SF = 6 dB.")
EQ(r"\mathrm{SF}_{\mathrm{dB}}=\max\left(0,\ \mu_{\mathrm{SF}}+\sigma_{\mathrm{SF}}Z\right),\qquad Z\sim\mathcal{N}(0,1)")
P("대류권 섬광 감쇠 (Table 6.6.6.2.1-1, 20 GHz, 1% 초과 확률값을 상한으로 균등 샘플링):")
EQ(r"A_{\mathrm{tropo,dB}}\sim\mathcal{U}\left[0,\ A_{1\%}(\varepsilon_e)\right]")
P("대규모 감쇠는 에피소드 내 상수 진폭으로 곱해진다. 이 상수는 5절의 인과적 정규화에서 소거되므로, "
  "예측 문제에는 소규모 페이딩과 위상 성분만 남는다.")
EQ(r"\gamma_{\mathrm{LS}}=10^{-(\mathrm{SF}_{\mathrm{dB}}+A_{\mathrm{tropo,dB}})/20}")
IMPL("losProbability, kaBandDelaySpreadParams, kaBandKFactorParams, kaBandShadowFadeParams, kaBandTropoAttenBoundDB")

H2("3.3 NTN-CDL 소규모 페이딩과 위성 배열 응답")
P("nrCDLChannel (NTN-CDL, SatelliteElevationAngle = ε_e, DelaySpread = DS, LOS이면 KFactor = K_dB)이 "
  "TR 38.901 CDL 절차와 TR 38.811 NTN-CDL 정의로 경로 이득을 만든다. 송신 소자 (m, n), 클러스터 p의 경로 이득 (개요):")
EQ(r"g_{p,mn}(t)=\sum_{r\in\mathcal{R}_p}\sqrt{\frac{P_p}{|\mathcal{R}_p|}}\,F_{\mathrm{tx}}(\Omega_{p,r})\,e^{j\frac{2\pi}{\lambda}\hat{\mathbf{r}}(\Omega_{p,r})^{T}\mathbf{d}_{mn}}\,e^{j\Phi_{p,r}}\,e^{j2\pi\nu_{p,r}t}")
EQ(r"\nu_{p,r}=f_{D,\mathrm{UE}}\cos\vartheta_{p,r}+f_{D,\mathrm{sat}}")
P("Ω: 위성 쪽 출발각, r̂: 그 방향 단위벡터, d_mn: 소자 위치, F_tx: 소자 패턴, P_p·τ_p: CDL 표의 정규화 지연에 DS를 곱한 값, "
  "Φ: 무작위 초기 위상, ϑ: 단말 이동 방향과 도착각 사이 각. LOS 클러스터 전력은 K-factor로 재배분. "
  "위성 Doppler는 MaximumDopplerShift = [f_D,UE, f_D,sat^max]로 주며 nrCDLChannel이 cos ε를 곱해 적용한다.")
EQ(r"f_{D,\mathrm{sat}}^{\max}=\frac{v_{\mathrm{sat}}}{c}\,\frac{R_e}{R_e+h}\,f_c\approx 466.7\ \mathrm{kHz},\qquad f_{D,\mathrm{sat}}(\varepsilon)=f_{D,\mathrm{sat}}^{\max}\cos\varepsilon")
P("위성 쪽 rank-1 구조: 위성에서 본 클러스터 각도 퍼짐이 사실상 0이므로 모든 경로가 같은 조향 벡터를 공유한다. "
  "a는 위성–단말 방향(에피소드 내 고정), s_q(i)는 페이딩·Doppler를 담은 시변 스칼라다. "
  "v9b 배열 자세(TransmitArrayOrientation = [0; 0; 0], 방위각 0°)에서 LOS 위상 기울기는 측정상 α = π sin ε, β = 0이다.")
EQ(r"\mathbf{h}_q(i)\approx s_q(i)\,\mathbf{a}(\alpha,\beta),\qquad [\mathbf{a}]_{mn}=\frac{1}{16}e^{j(\alpha m+\beta n)}")

H2("3.4 OFDM 주파수 응답과 관측")
P("관측 i마다 시각 tᵢ의 슬롯 하나만 생성하고(InitialTime = tᵢ), 경로 이득에 4절의 위상 항 ψ(t)를 곱한 뒤 완전 채널 추정으로 "
  "주파수 응답을 얻는다. τ₀는 에피소드 첫 관측에서 한 번 추정해 고정한 타이밍 오프셋(v9b)이다. "
  "고정 오프셋이므로 관측 간 부반송파 위상 기울기 점프가 없다 (로더의 점프 보정은 v9b 전 split에서 0건 작동). "
  "아래는 경로 필터의 대역 제한을 무시한 근사식이며 구현은 nrPerfectChannelEstimate의 정확한 계산이다.")
EQ(r"[\mathbf{h}_q(i)]_{mn}\approx\sum_{p}g_{p,mn}(t_i)\,\psi(t_i)\,e^{-j2\pi(q-150)\Delta f\,(\tau_p-\tau_0)},\qquad \Delta f=60\ \mathrm{kHz}")
IMPL("관측 루프, nrPerfectTimingEstimate (첫 관측 1회), nrPerfectChannelEstimate(..., offset, sampleTimes), H(:, end, :, :)")

H1("4. 비선형·비정상 성분 모델링")
P("경로 이득에 공통으로 곱하는 위상 항 (실제 샘플 시각 t 기준). 첫째 항은 궤도 Doppler 사전보상, 둘째 항은 Doppler 변화율 잔차, "
  "셋째 항은 LO 위상잡음이다.")
EQ(r"\psi(t)=\exp\left(-j2\pi f_{D,\mathrm{sat}}\,t+j\pi\dot{f}_D\,t^2+j\varphi_{\mathrm{PN}}(t)\right)")
H2("4.1 궤도 Doppler 사전보상")
P("궤도 정보(ephemeris) 기반 개루프 보상으로 위성 Doppler f_D,sat를 정확히 상쇄한다. "
  "남는 시변 성분은 단말 Doppler, Doppler 변화율 잔차, LO 위상잡음이다.")
H2("4.2 Doppler 변화율 잔차 (2차 위상)")
P("TR 38.811은 위성 속도·고각을 ‘few TTI’ 동안만 상수로 볼 수 있다고 명시한다. 150 ms 에피소드에서는 Doppler의 시간 변화가 남는다. "
  "원궤도·구형 지구에서 위성 직하점–단말 지구 중심각 θ에 대해:")
EQ(r"d(\theta)=\sqrt{R_e^2+R_h^2-2R_eR_h\cos\theta},\qquad f_D(\theta)=-\frac{f_c}{c}\,\omega\,\frac{R_eR_h\sin\theta}{d(\theta)}")
EQ(r"\cos\theta_0(\varepsilon)=x\cos^2\varepsilon+\sin\varepsilon\sqrt{1-x^2\cos^2\varepsilon},\qquad x=\frac{R_e}{R_h},\quad \omega=\frac{v_{\mathrm{sat}}}{R_h},\quad R_h=R_e+h")
EQ(r"\dot{f}_D^{\mathrm{raw}}(\varepsilon)=\omega^2\frac{\partial}{\partial\theta}\left[-\frac{f_c}{c}\frac{R_eR_h\sin\theta}{d(\theta)}\right]_{\theta=\theta_0(\varepsilon)}")
P("미분은 Δθ = 10⁻⁵ rad 중앙 차분이다. 샘플링 Nyquist(1/(2Δt) = 500 Hz) 여유를 위해 에피소드 말미 잔차 순간주파수가 150 Hz를 넘지 않도록, "
  "최악 조건 ε = 90° 기준 전역 스케일을 한 번만 정해 모든 에피소드에 같은 비율로 적용한다 (고각별 상대 차이 보존, f_D^raw(90°) = −6449.3 Hz/s). "
  "π ḟ_D t²는 선형 위상 램프(상수 회전)로 제거되지 않는 2차 위상이다.")
EQ(r"\kappa=\min\left(1,\ \frac{150\ \mathrm{Hz}}{|\dot{f}_D^{\mathrm{raw}}(90^\circ)|\cdot 0.15\ \mathrm{s}}\right)=0.1551,\qquad \dot{f}_D=\kappa\,\dot{f}_D^{\mathrm{raw}}(\varepsilon_e)")
IMPL("dopplerRateAtElev, elevToCentralAngle, dopplerCentralAngleFunc, globalDopplerRateScale")
H2("4.3 LO 위상잡음 (Wiener 과정)")
P("관측마다 한 번 갱신하고 그 슬롯 안에서는 상수다 (150 ms 누적 표준편차 약 3.7°).")
EQ(r"\varphi_{\mathrm{PN}}(t_i)=\varphi_{\mathrm{PN}}(t_{i-1})+\sigma_{\mathrm{PN}}\sqrt{\Delta t}\,w_i,\qquad w_i\sim\mathcal{N}(0,1),\quad \sigma_{\mathrm{PN}}=0.3^\circ/\sqrt{\mathrm{ms}}")
IMPL("phiAccum = phiAccum + phaseNoiseStdRadPerSqrtMs·sqrt(Δt)·randn,  ph = exp(-1i*2*pi*fD_sat*ts + 1i*pi*fDdot_sat*ts.^2 + 1i*phiAccum)")
H2("4.4 저장")
P("CSI 실수·허수부를 에피소드별 최대 크기로 int16 양자화해 저장하고 (SQNR 약 94 dB, 관측 SNR보다 충분히 큼), "
  "인과적 정규화용 누적 최댓값을 함께 저장한다 (300 부반송파·256 안테나 전체 최대).")
EQ(r"c_e(i)=\max_{i'\le i}\ \max_{q,m,n}\left|[\mathbf{H}_q(i')]_{mn}\right|")
IMPL("dsReal, dsImag, dsScale, causalMaxPerSlot = cummax(perSlotMax)")

H1("5. 학습 샘플 구성 (온더플라이 로더)")
H2("5.1 윈도우")
P("에피소드 e, 시작 w, 부반송파 q ∈ 𝒬마다 샘플 하나: history i = w…w+T_p−1, 미래 i = w+T_p…w+T_p+T_f−1, "
  "예측 대상 i* = w+T_p−1+k. w = 0…150−T_p−T_f (stride 1) → 에피소드당 115 윈도우 (T_f = 16) × 12 부반송파.")
H2("5.2 인과적 정규화")
EQ(r"\bar{\mathbf{h}}_q(i)=\frac{\mathbf{h}_q(i)}{c_e(w+T_p-1)}")
H2("5.3 채널 추정 잡음")
P("SNR은 윈도우마다 U[5, 20] dB에서 뽑고 12 부반송파에 공통이다. 잡음은 history에만 넣고 정답은 잡음 없는 h̄다. "
  "평가 배치는 고정 시드로 만들어 모든 방법이 같은 윈도우·같은 잡음을 쓴다.")
EQ(r"\tilde{\mathbf{h}}_q(i)=\bar{\mathbf{h}}_q(i)+\mathbf{n}_q(i),\qquad \mathbf{n}_q(i)\sim\mathcal{CN}\left(\mathbf{0},\sigma_n^2\mathbf{I}\right),\qquad \sigma_n^2=P_h\,10^{-\mathrm{SNR}/10}")
EQ(r"P_h=\frac{1}{|\mathcal{Q}|\,T_p\,256}\sum_{q\in\mathcal{Q}}\sum_{i=w}^{w+T_p-1}\left\|\bar{\mathbf{h}}_q(i)\right\|^2,\qquad \mathrm{SNR}\sim\mathcal{U}[5,20]\ \mathrm{dB}")
H2("5.4 Doppler 잔차 회전 제거 (인과적)")
P("history(o = −(T_p−1)…0)와 미래(o = 1…T_f), 잡음 관측과 정답에 같은 회전을 곱해 예측 관계를 보존한다. 이하 h̃, h̄는 회전 제거 후 신호다.")
EQ(r"\hat{\omega}=\angle\sum_{q\in\mathcal{Q}}\sum_{i=w+1}^{w+T_p-1}\tilde{\mathbf{h}}_q(i-1)^{H}\tilde{\mathbf{h}}_q(i)")
EQ(r"\mathbf{x}(i)\leftarrow\mathbf{x}(i)\,e^{-j\hat{\omega}\,o_i},\qquad o_i=i-(w+T_p-1)")
H2("5.5 학습 증강")
P("학습 배치에만 샘플별 공통 위상 회전을 입력 전 프레임과 정답에 곱한다 (절대 위상은 정보가 아님).")
EQ(r"\left(\tilde{\mathbf{h}},\bar{\mathbf{h}}\right)\leftarrow e^{j\phi}\left(\tilde{\mathbf{h}},\bar{\mathbf{h}}\right),\qquad \phi\sim\mathcal{U}[0,2\pi)")
IMPL("data.Windows.batch, data.to_model_io, data.phase_augment")

H1("6. 예측 문제와 평가 지표")
P("모델 입력은 2T_p×16×16 실수 텐서 (Re H̃(w), Im H̃(w), … 시간순), 출력은 2×16×16이다.")
EQ(r"\hat{\mathbf{h}}_q(i^\star)=f_\Theta\left(\tilde{\mathbf{h}}_q(w),\dots,\tilde{\mathbf{h}}_q(w+T_p-1)\right),\qquad i^\star=w+T_p-1+k")
P("NMSE는 윈도우·부반송파·256 안테나 전체를 합산(pooled)해 계산한다 (그룹별: 전체, LOS, NLOS).")
EQ(r"\mathrm{NMSE}_k=10\log_{10}\frac{\sum_{(e,w)}\sum_{q\in\mathcal{Q}}\left\|\hat{\mathbf{h}}_q(i^\star)-\bar{\mathbf{h}}_q(i^\star)\right\|^2}{\sum_{(e,w)}\sum_{q\in\mathcal{Q}}\left\|\bar{\mathbf{h}}_q(i^\star)\right\|^2}\ [\mathrm{dB}]")
IMPL("evaluate.py")

H1("7. 제안 모델: BSC-UNet (Beam-ShortCut UNet)")
P("전체 구조: 첫째 항은 빔 지름길(물리 구조: 추정 빔 방향 × 예측 빔 계수), 둘째 항은 UNet 보정(학습)이다.")
EQ(r"\hat{\mathbf{h}}(i^\star)=\hat{s}(i^\star)\,\hat{\mathbf{a}}+U_\Theta\left(\left[\tilde{\mathbf{H}},\,\mathbf{S}\right]\right)")
H2("7.1 빔 방향 추정 (학습 파라미터 없음)")
P("rank-1 구조의 방향을 history 전체 에너지 최대화로 찾는다.")
EQ(r"(\hat{\alpha},\hat{\beta})=\arg\max_{\alpha,\beta}\sum_{\tau=w}^{w+T_p-1}\left|\mathbf{a}(\alpha,\beta)^{H}\tilde{\mathbf{h}}(\tau)\right|^2")
P("① 거친 탐색: 64×64 zero-pad 2D FFT.")
EQ(r"P(k_1,k_2)=\sum_{\tau}\left|\sum_{m,n}[\tilde{\mathbf{H}}(\tau)]_{mn}\,e^{-j2\pi(k_1m+k_2n)/64}\right|^2,\qquad (\alpha_0,\beta_0)=\frac{2\pi}{64}\left(k_1^\star,k_2^\star\right)")
P("② 정밀 탐색: α, β 각각 ±1 bin을 65점으로 나눠 같은 목적함수의 최대점. ③ 단위 노름 조향 벡터 구성.")
EQ(r"\alpha\in\alpha_0+\frac{2\pi}{64}\{-1,\dots,1\},\qquad [\hat{\mathbf{a}}]_{mn}=\frac{1}{16}e^{j(\hat{\alpha}m+\hat{\beta}n)}")
IMPL("csi_models.manifold_beam (NF = 64, N_FINE = 65)")
H2("7.2 빔 지름길 (학습형 선형 시간 예측)")
P("복소 T_p탭 필터 w_l은 0으로 초기화한다.")
EQ(r"s(\tau)=\hat{\mathbf{a}}^{H}\tilde{\mathbf{h}}(\tau),\qquad \hat{s}(i^\star)=\sum_{l=0}^{T_p-1}w_l\,s(w+l),\qquad w_l\in\mathbb{C}")
IMPL("_beam_shortcut, 파라미터 w_re, w_im")
H2("7.3 UNet 보정")
P("입력은 원래 history와 빔 투영(잡음 제거) history의 연결이다 (4T_p = 80 채널).")
EQ(r"\mathbf{S}(\tau)=s(\tau)\,\hat{\mathbf{a}},\qquad U_\Theta:\mathbb{R}^{4T_p\times16\times16}\rightarrow\mathbb{R}^{2\times16\times16}")
TB(["층", "연산", "크기"], [
    ["down1", "Conv 4×4 / stride 2, LeakyReLU(0.2)", "80×16² → 64×8²"],
    ["down2", "Conv 4×4 / s2, InstanceNorm, LeakyReLU", "→ 128×4²"],
    ["down3", "Conv 4×4 / s2, InstanceNorm, LeakyReLU, Dropout 0.5", "→ 256×2²"],
    ["down4", "Conv 4×4 / s2, LeakyReLU, Dropout 0.5", "→ 512×1²"],
    ["bottleneck", "3 × [1×1 Conv, 채널 정규화, LeakyReLU, Dropout] + 잔차 연결", "512×1²"],
    ["up1–up3", "ConvTranspose 4×4 / s2, Norm, ReLU (+Dropout), skip concat", "→ 128×8²"],
    ["final", "Upsample ×2, pad, Conv 4×4", "→ 2×16²"],
])
IMPL("csi_models.GeneratorUNet, BeamShortcutUNet.forward")
H2("7.4 학습")
P("빔 필터 w와 UNet Θ를 end-to-end로 함께 학습한다 (MSE, NMSE와 일치).")
EQ(r"\mathcal{L}(\Theta,\mathbf{w})=\frac{1}{N}\sum_{\text{batch}}\left\|\hat{\mathbf{h}}(i^\star)-\bar{\mathbf{h}}(i^\star)\right\|^2")
TB(["항목", "값"], [
    ["옵티마이저", "Adam, η₀ = 10⁻⁴, (β₁, β₂) = (0.9, 0.999), gradient clipping 1.0"],
    ["학습률", "cosine, 80 epoch 동안 10⁻⁴ → 10⁻⁶"],
    ["가중치 평균", "EMA 0.999 (검증·저장·평가에 EMA 가중치 사용)"],
    ["배치", "윈도우 4개 × 부반송파 12 = 샘플 48개"],
    ["선택·종료", "3-epoch 이동평균 val NMSE 최소 저장, patience 20, 최대 100 epoch, seed 42"],
])
IMPL("train_beam.py --lr_sched cosine --cosine_epochs 80 --ema 0.999 --beta1 0.9 --grad_clip 1.0 --patience 20 --epochs 100")

H1("8. 비교 기법")
P("모든 비교 기법은 같은 입력(잡음 history T_p개, 부반송파별 독립)만 쓴다. 학습형 기법은 train으로 적합하고 정규화 계수 α는 "
  "{0, 10⁻⁵, 10⁻⁴, 10⁻³, 10⁻², 10⁻¹, 1}에서 val로 horizon마다 고른다. 시간 자기상관(train clean, 구간 길이 T_p+k, "
  "안테나·부반송파 풀링)과 ridge는 다음과 같다. history를 시간순 벡터 x̃ = [x̃(w), …, x̃(w+T_p−1)]ᵀ로 둔다.")
EQ(r"r_k(\tau)=\mathbb{E}\left[x(t)\,x^*(t+\tau)\right],\qquad \mathrm{ridge}(\mathbf{A},\mathbf{B};\alpha)=\left(\mathbf{A}+\lambda\mathbf{I}\right)^{-1}\mathbf{B},\quad \lambda=\alpha\frac{\mathrm{tr}\,\mathbf{A}}{T_p}")
P("Persistence (회전 제거 후이므로 위상 선형 외삽과 같음):")
EQ(r"\hat{\mathbf{h}}(i^\star)=\tilde{\mathbf{h}}(w+T_p-1)")
P("LMMSE (Wiener, 안테나별 스칼라 필터, σ_n²는 train 평균 잡음 분산):")
EQ(r"\hat{x}(i^\star)=\mathbf{w}_k^{H}\tilde{\mathbf{x}},\qquad \mathbf{w}_k=\left(\mathbf{R}_k+\sigma_n^2\mathbf{I}\right)^{-1}\mathbf{r}_k,\qquad [\mathbf{R}_k]_{ab}=r_k(b-a),\quad [\mathbf{r}_k]_a=r_k(T_p-1+k-a)")
P("Kalman (AR(1) 상태 x(i+1) = φx(i) + u(i), 관측 x̃ = x + n, P₀ = r_k(0)):")
EQ(r"\varphi=\frac{r_k(1)^*}{r_k(0)},\qquad K_l=\frac{P_l}{P_l+\sigma_n^2},\qquad P_{l+1}=|\varphi|^2(1-K_l)P_l+r_k(0)\left(1-|\varphi|^2\right),\qquad \hat{x}(i^\star)=\varphi^k\,\hat{x}_{\mathrm{post}}(w+T_p-1)")
P("LinReg_eig (train 공간 공분산 Σ = E[hhᴴ]의 고유기저 U = [u₁, …, u₂₅₆]에서 방향별 ridge):")
EQ(r"c_j(\tau)=\mathbf{u}_j^{H}\tilde{\mathbf{h}}(\tau),\qquad \hat{c}_j=\mathbf{c}_j^{T}\mathbf{W}_{j,k},\qquad \hat{\mathbf{h}}=\sum_{j=1}^{256}\hat{c}_j\,\mathbf{u}_j")
P("Poly2_blind (학습 없음): 고유빔 g = history 행렬의 최대 특이벡터 방향, s(τ) = gᴴh̃(τ). 언랩 위상과 크기를 "
  "o ∈ {−(T_p−1), …, 0}에서 최소제곱 적합해 o = k로 외삽한다.")
EQ(r"\arg s(o)\approx c_2o^2+c_1o+c_0,\qquad |s(o)|\approx d_1o+d_0")
EQ(r"\hat{\mathbf{h}}(i^\star)=\max\left(0,\,d_1k+d_0\right)e^{j(c_2k^2+c_1k+c_0)}\,\mathbf{g}")
P("BeamLinReg: 같은 고유빔 g, 미래 빔 계수 gᴴh̄(i*)를 목표로 train ridge 적합.")
EQ(r"\hat{s}=\mathbf{s}^{T}\mathbf{w}_k,\qquad \hat{\mathbf{h}}(i^\star)=\hat{s}\,\mathbf{g}")
P("ManifoldLinReg: BeamLinReg에서 g ← â (배열 매니폴드 빔, manifold_beam). BSC-UNet에서 UNet을 뺀 구조이며 지름길 필터를 닫힌 식으로 적합한다. "
  "ManifoldPoly2 (ablation): Poly2_blind에서 g ← â. 매니폴드 빔 + 2차 위상 외삽 (학습 없음).")
IMPL("baselines.Baselines.fit / predict, _blind_beam, _manifold_v_fn, _ridge, _autocorr")

H1("9. 통계 검정과 누수 검사")
P("에피소드 bootstrap: 에피소드 e별 오차 합 N_e와 에너지 합 D_e로 Δ를 계산하고, 에피소드를 복원추출 10,000회 해서 "
  "2.5–97.5% 구간을 95% 신뢰구간으로 쓴다 (Δ < 0: A 우세, 구간이 0을 포함하지 않으면 유의). 에피소드별 승패 수도 보고한다.")
EQ(r"\Delta=10\log_{10}\frac{\sum_{e}N_e^{(A)}}{\sum_{e}D_e}-10\log_{10}\frac{\sum_{e}N_e^{(B)}}{\sum_{e}D_e}")
P("누수 검사: (T2) 정답을 윈도우 간 섞었을 때 NMSE ≥ 0 dB, (T3) 출력과 k′ 시점 정답의 NMSE가 학습한 k에서 최소.")
IMPL("evaluate.py --leak_check")

H1("10. 확장: 기하 일관 데이터(v10)와 빔 정렬 모델(BSC-UNet-AI)")
H2("10.1 구면 궤도 기하")
P("에피소드마다 고각 ε_e(같은 스윕·split)에 대해 단말 cross-track 중심각 ψ와 통과 방향 σ를 뽑고 (전용 난수열), 지구 중심 좌표에서 위성·단말 위치와 속도를 정한다.")
EQ(r"\psi\sim\mathcal{U}[-\gamma,\gamma],\quad \gamma=\theta_0(\varepsilon_e),\quad \sigma\in\{\pm1\},\qquad \theta=\sigma\arccos\frac{\cos\gamma}{\cos\psi}")
EQ(r"\mathbf{r}_{\mathrm{sat}}=R_h[\sin\theta,\,0,\,\cos\theta]^{T},\qquad \mathbf{r}_{\mathrm{UE}}=R_e[0,\,\sin\psi,\,\cos\psi]^{T},\qquad \mathbf{v}=v_{\mathrm{sat}}[\cos\theta,\,0,\,-\sin\theta]^{T}")
P("l = r_sat − r_UE, 위성 본체 축 x_b = v/v_sat, z_b = −r_sat/R_h (nadir), y_b = z_b × x_b, u = −l/‖l‖. "
  "생성기는 ḟ_D를 1 ms 중앙 차분으로 계산하며 아래 닫힌 식과 상대오차 10⁻⁹ 미만으로 일치한다.")
EQ(r"f_{D,\mathrm{geo}}=-\frac{f_c}{c}\,\frac{\mathbf{v}^{T}\mathbf{l}}{\|\mathbf{l}\|},\qquad \dot{f}_{D,\mathrm{geo}}=-\frac{f_c}{c}\,\frac{\|\mathbf{v}\|^2+\mathbf{l}^{T}\mathbf{a}_c-\left(\mathbf{v}^{T}\mathbf{l}/\|\mathbf{l}\|\right)^2}{\|\mathbf{l}\|},\qquad \mathbf{a}_c=-\omega^2\mathbf{r}_{\mathrm{sat}}")
EQ(r"\phi_{\mathrm{sat}}=\mathrm{atan2}\left(\mathbf{u}^{T}\mathbf{y}_b,\ \mathbf{u}^{T}\mathbf{x}_b\right),\qquad \eta=\arccos\left(\mathbf{u}^{T}\mathbf{z}_b\right),\qquad \sin\eta=\frac{R_e}{R_h}\cos\varepsilon")
P("nrCDLChannel에는 MaximumDopplerShift(2) = |f_D,geo|/cos ε (적용값 = |f_D,geo|), Doppler 변화율 잔차에는 κ·ḟ_D,geo를 쓴다. "
  "nadir 지향 배열을 단말 지점 좌표계에서 표현하는 회전(TR 38.901 LCS–GCS 변환, bearing–downtilt–slant)은 다음과 같다. "
  "이때 LOS 위상 기울기의 크기는 π sin η, 방향은 180° − φ_sat이다 (생성 데이터로 확인). "
  "TR 38.811 통계 파라미터(K, DS, SF)는 고각만의 함수로 그대로 쓰고, 방위각은 배열 공간 응답에만 들어간다.")
EQ(r"\mathrm{Orientation}=\left[\,0;\ 90^\circ-\gamma_c;\ \phi_{\mathrm{sat}}\,\right],\qquad \gamma_c=90^\circ-\varepsilon-\eta,\qquad \sqrt{\alpha^2+\beta^2}=\pi\sin\eta")
IMPL("generate_ntn_csi_v10_fdd_dl.m: passGeometry, passState; 검증 check_doppler_rate_v10.py, validate_v10.py")
H2("10.2 빔 정렬 모델 (BSC-UNet-AI)")
P("빔 지름길은 7절과 같고, UNet은 추정 방향을 제거한 정렬 좌표에서 잔차를 예측한 뒤 방향을 되돌린다 (replicate 패딩). "
  "방향이 달라도 UNet이 보는 입력이 같아지므로 방위각에 대한 일반화가 구조적으로 확보된다.")
EQ(r"\mathbf{u}=16\,\hat{\mathbf{a}},\qquad \tilde{\mathbf{H}}^{\mathrm{al}}(\tau)=\tilde{\mathbf{H}}(\tau)\odot\mathbf{u}^{*},\qquad \mathbf{S}^{\mathrm{al}}(\tau)=\frac{s(\tau)}{16}\,\mathbf{1}_{16\times16}")
EQ(r"\hat{\mathbf{h}}(i^\star)=\hat{s}(i^\star)\,\hat{\mathbf{a}}+U_\Theta\left(\left[\tilde{\mathbf{H}}^{\mathrm{al}},\,\mathbf{S}^{\mathrm{al}}\right]\right)\odot\mathbf{u}")
IMPL("csi_models.BeamAlignedUNet, train_beam.py --aligned")
H2("10.3 방위각 일반화 시험 (v9b 모델)")
P("다른 방위각의 단말은 열 방향 위상 기울기를 곱해 모사한다 (rank-1 구조에서 β 변화와 등가).")
EQ(r"[\mathbf{H}^{(\beta)}_q(i)]_{mn}=[\mathbf{H}_q(i)]_{mn}\,e^{j\beta n},\qquad \beta\in\{0,\,0.5,\,1,\,2,\,3\}\ \mathrm{rad}")
IMPL("eval_bearing.py")

H1("11. 기호–구현 대응표")
TB(["수식", "구현"], [
    ["ε_e, split", "generateElevationProfile, testBandsDeg, valBandsDeg, testGuardsDeg"],
    ["P_LOS, DS, K, SF, A_tropo", "losProbability, kaBand*Params, kaBandTropoAttenBoundDB"],
    ["f_D,sat, ḟ_D, κ", "fD_sat, fDdot_sat, globalDopplerRateScale"],
    ["ψ(t)", "ph = exp(-1i*2*pi*fD_sat*ts + 1i*pi*fDdot_sat*ts.^2 + 1i*phiAccum)"],
    ["τ₀", "offset = nrPerfectTimingEstimate(...) (첫 관측 1회)"],
    ["c_e(i)", "causalMaxPerSlot → data.Split.cm"],
    ["σ_n², SNR, ω̂", "data.Windows.batch: std, snr, omega"],
    ["â", "csi_models.manifold_beam"],
    ["s, ŝ, w_l", "_beam_shortcut, w_re, w_im"],
    ["U_Θ", "csi_models.GeneratorUNet"],
    ["BSC-UNet / BSC-UNet-AI", "BeamShortcutUNet / BeamAlignedUNet"],
    ["NMSE, bootstrap", "evaluate.py (groups, paired)"],
])


# ------------------------------------------------------------------ Word 작성
def set_font(run, size=None, mono=False, color=None, bold=None):
    run.font.name = "Consolas" if mono else "Malgun Gothic"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = parse_xml(r'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size: run.font.size = Pt(size)
    if color: run.font.color.rgb = RGBColor(*color)
    if bold is not None: run.font.bold = bold


def word_latex(latex):
    """Word 수식 변환용 전처리 (Office MML2OMML.XSL이 mspace를 버리므로 간격을 공백 문자로, argmax는 아래첨자 형태로).
    문서에 표시하는 LaTeX 코드는 원본(표준 문법) 그대로 둔다."""
    t = latex.replace(r"\arg\max_{\alpha,\beta}", r"\underset{\alpha,\beta}{\mathrm{arg\,max}}")
    for a, b in ((r"\qquad", "~~~~~~"), (r"\quad", "~~~"), (r"\ ", "~"), (r"\;", "~"), (r"\,", "~")):
        t = t.replace(a, b)
    return t


def add_equation(doc, latex):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        mml = tex2mml(word_latex(latex))
        omml = _xslt(etree.fromstring(mml.encode("utf-8")))
        p._element.append(parse_xml(etree.tostring(omml.getroot())))
    except Exception as e:                                    # 변환 실패 시 LaTeX 코드만 남김
        FAILED.append((latex, repr(e)[:120]))
        set_font(p.add_run("[Word 수식 변환 실패: 아래 LaTeX 코드 사용]"), 9, color=(192, 0, 0))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.LEFT
    c.paragraph_format.space_after = Pt(8)
    set_font(c.add_run(latex), 8.5, mono=True, color=(110, 110, 110))


def add_table(doc, hdr, rows):
    t = doc.add_table(rows=1, cols=len(hdr)); t.style = "Table Grid"
    for j, h in enumerate(hdr):
        cell = t.rows[0].cells[j]; cell.text = ""
        set_font(cell.paragraphs[0].add_run(h), 9.5, bold=True)
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r):
            cells[j].text = ""
            set_font(cells[j].paragraphs[0].add_run(v), 9.5)
    doc.add_paragraph()


def build():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Malgun Gothic"; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    for name in ("Title", "Heading 1", "Heading 2"):
        s = doc.styles[name]; s.font.name = "Malgun Gothic"
        s.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    doc.add_heading("LEO NTN Ka-band FDD 하향링크 CSI 예측: 데이터 생성 · BSC-UNet · 비교 기법 · 평가 수식 정리", level=0)
    for item in C:
        kind = item[0]
        if kind == "h1":
            doc.add_heading(item[1], level=1)
        elif kind == "h2":
            doc.add_heading(item[1], level=2)
        elif kind == "p":
            set_font(doc.add_paragraph().add_run(item[1]), 10.5)
        elif kind == "eq":
            add_equation(doc, item[1])
        elif kind == "impl":
            p = doc.add_paragraph()
            set_font(p.add_run("구현: "), 9, bold=True, color=(40, 90, 160))
            set_font(p.add_run(item[1]), 9, mono=True, color=(40, 90, 160))
        elif kind == "table":
            add_table(doc, item[1], item[2])
    doc.save(OUT)
    n_eq = sum(1 for c in C if c[0] == "eq")
    print(f"saved {OUT}  equations={n_eq}  Word 수식 변환 실패={len(FAILED)}")
    for lt, err in FAILED:
        print("  FAIL:", lt[:90], err)


if __name__ == "__main__":
    build()
